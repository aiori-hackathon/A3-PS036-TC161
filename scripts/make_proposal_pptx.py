"""Build docs/proposal/A3-PS036-TC161.pptx following the AIORI Hackathon 3 proposed structure.

Usage: pip install python-pptx==1.0.2 && python scripts/make_proposal_pptx.py
Numbers on the optimization slide come from results/runs/20261010T130246Z-ps036-pinned (with live run).
"""

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "proposal" / "A3-PS036-TC161.pptx"
LOGO = ROOT / "docs" / "proposal" / "aiori-logo.png"
ARCH = ROOT / "docs" / "proposal" / "solution-architecture.png"

TEAM = "A3-PS036-TC161"
AUTHOR, EMAIL = "Akash Joshi", "akashjoshi2023@gmail.com"
MEMBERS = [("Akash Joshi", "akashjoshi2023@gmail.com"),
           ("Anitha", "anithapriyanatarajan@gmail.com"),
           ("Induradhakrishnan", "induradhakrishnan@pes.edu")]
EVENT, DATE = "AIORI Hackathon 3", "October 7, 2026"
TITLE = "PQC DNSSEC MTL Benchmark Using a Reference Implementation (PS-036)"

BLUE, BLACK, GREY, WHITE = "0000DC", "000000", "555555", "FFFFFF"
PALE, LIGHT = "F2F2FD", "E3E3FA"
GREEN, AMBER = "2E7D32", "B26A00"
FONT = "Calibri"
MARGIN, WIDTH = 0.59, 12.15
TOKEN = re.compile(r"(\*\*.+?\*\*|~~.+?~~)")


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rich(p, text: str, size: float, color: str = BLACK, bold: bool = False) -> None:
    """Add runs to a paragraph; **bold** and ~~italic~~ spans are honoured."""
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i = bold, False
        if part.startswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("~~"):
            part, i = part[2:-2], True
        r = p.add_run()
        r.text = part
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), b, i
        f.color.rgb = rgb(color)


def hanging(p, indent: float, char: str | None = "•") -> None:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(indent)))
    pPr.set("indent", str(-Inches(indent)))
    if char is None:
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        return
    clr = pPr.makeelement(qn("a:buClr"), {})
    clr.append(clr.makeelement(qn("a:srgbClr"), {"val": BLUE}))
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}))


def textbox(slide, x, y, w, h, paras, size=16, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            bold=False, space=0, bullets=False, indent=0.3):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for k, text in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        if k and space:
            p.space_before = Pt(space)
        rich(p, text, size, color, bold)
        if bullets:
            hanging(p, indent)
    return tb


def box(slide, x, y, w, h, lines, fill=WHITE, line=BLUE, size=11, shape=MSO_SHAPE.RECTANGLE, dash=False,
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, lw=1.25):
    """lines: str or (text, size[, color]) per paragraph."""
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for k, item in enumerate([lines] if isinstance(lines, str) else lines):
        text, sz, color = (item, size, BLACK) if isinstance(item, str) else (*item, BLACK)[:3]
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        rich(p, text, sz, color)
    return sh


def arrow(slide, x1, y1, x2, y2, color=BLACK, dash=False, both=False, head=True, width=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = c.line._get_or_add_ln()
    if both:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def table(slide, x, y, widths, rows, size=12, row_h=0.36):
    gf = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)),
                                Inches(row_h * len(rows)))
    tbl = gf.table
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(BLUE if i == 0 else (PALE if i % 2 else WHITE))
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            for k, text in enumerate(value.split("\n")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                rich(p, text, size, WHITE if i == 0 else BLACK, bold=i == 0)
    return gf


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.framed = []

    def blank(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def frame(self, section, title, lead=None, root=True):
        s = self.blank()
        if root:
            textbox(s, 3.2, 0.08, 3.5, 0.3, "AIORI 3", 12, BLUE, PP_ALIGN.RIGHT, bold=True)
            textbox(s, 6.85, 0.08, 6.0, 0.3, section, 12, BLUE, bold=True)
        else:
            textbox(s, 0, 0.08, 13.333, 0.3, section, 12, BLUE, PP_ALIGN.CENTER, bold=True)
        textbox(s, MARGIN, 0.42, WIDTH, 0.6, title, 30, BLUE, bold=True)
        if lead:
            textbox(s, MARGIN, 1.07, WIDTH, 0.45, lead, 17, bold=True)
        self.framed.append(s.slide_id)
        return s

    def caption(self, s, text, y=6.42):
        textbox(s, 0, y, 13.333, 0.35, f"Figure: {text}", 15, align=PP_ALIGN.CENTER)
        p = s.shapes[-1].text_frame.paragraphs[0]
        p.runs[0].text = "Figure:"
        p.runs[0].font.color.rgb = rgb(BLUE)
        r = p.add_run()
        r.text = f" {text}"
        r.font.name, r.font.size = FONT, Pt(15)
        r.font.color.rgb = rgb(BLACK)

    def save(self, path: Path):
        ids = [s.slide_id for s in self.prs.slides]
        for s in self.prs.slides:
            if s.slide_id not in self.framed:
                continue
            textbox(s, MARGIN, 7.12, 8.0, 0.3, f"**{TEAM}**  •  {DATE}", 12, BLUE)
            textbox(s, 10.5, 7.12, 2.25, 0.3, f"{ids.index(s.slide_id) + 1} / {len(ids)}", 12, BLUE,
                    PP_ALIGN.RIGHT)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)


def title_slide(d: Deck):
    s = d.blank()
    s.shapes.add_picture(str(LOGO), Inches(0.58), Inches(0.58), Inches(1.18), Inches(1.18))
    textbox(s, MARGIN, 2.95, WIDTH, 0.75, TEAM, 44, BLUE, bold=True)
    textbox(s, MARGIN, 3.82, WIDTH, 0.5, TITLE, 25)
    textbox(s, MARGIN, 4.6, WIDTH, 1.3, [f"**{name}**:  {email}" for name, email in MEMBERS], 19, space=4)
    textbox(s, MARGIN, 6.02, WIDTH, 0.35, EVENT, 17)
    textbox(s, MARGIN, 6.55, WIDTH, 0.4, DATE, 19)


def flow(slide, x, y, w, h, steps, gap=0.45, vertical=False, size=15):
    """Boxes joined by arrows; the last box is highlighted."""
    for k, text in enumerate(steps):
        last = k == len(steps) - 1
        bx, by = (x, y + k * (h + gap)) if vertical else (x + k * (w + gap), y)
        box(slide, bx, by, w, h, [(f"**{text}**" if last else text, size)], fill=LIGHT if last else PALE,
            line=BLUE, lw=2 if last else 1.25)
        if not last:
            if vertical:
                arrow(slide, bx + w / 2, by + h, bx + w / 2, by + h + gap, BLUE, width=2)
            else:
                arrow(slide, bx + w, by + h / 2, bx + w + gap, by + h / 2, BLUE, width=2)


def problem_slide(d: Deck):
    s = d.frame("Abstract", "The Problem: DNSSEC in the Quantum Era", "What is at risk?")
    textbox(s, MARGIN, 1.75, 7.4, 3.3, [
        "**DNS** translates domain names into IP addresses so users can reach websites.",
        "**DNSSEC** uses digital signatures to verify DNS data and protect against tampering [1].",
        "**Quantum threat:** future quantum computers could break the classical cryptographic algorithms "
        "used by DNSSEC (RSA, ECDSA).",
        "**PQC** (Post-Quantum Cryptography) offers algorithms designed to resist quantum attacks, but some "
        "produce much larger keys and signatures [4].",
    ], 17, space=11, bullets=True)
    flow(s, 8.55, 1.8, 4.2, 0.78, ["Quantum-security requirements increase", "Cryptographic data becomes larger",
                                   "DNS response-size challenges emerge"], gap=0.42, vertical=True, size=15)

    textbox(s, MARGIN, 5.2, WIDTH, 0.35, "How one lookup works today", 15, BLUE, bold=True)
    steps = [("You type", "www.example.com"), ("Resolver", "asks the domain's DNS server"),
             ("DNS server", "answers: IP address + signature"), ("Resolver", "checks the signature, then connects")]
    w, gap = 2.6, 0.58
    for k, (head, text) in enumerate(steps):
        x = MARGIN + k * (w + gap)
        box(s, x, 5.62, w, 0.82, [(f"**{head}**", 14), (text, 12, GREY)], fill=WHITE)
        if k < len(steps) - 1:
            arrow(s, x + w, 6.03, x + w + gap, 6.03, BLUE, width=2)
    textbox(s, MARGIN, 6.55, WIDTH, 0.35, "The signature travels inside every answer, so its size is paid on "
            "every lookup.", 13, GREY)


def challenge_slide(d: Deck):
    s = d.frame("Abstract", "The Technical Challenge", "Why does larger cryptographic data matter?")
    textbox(s, MARGIN, 1.7, WIDTH, 0.8, "DNS commonly uses **UDP** for responses. Larger DNSSEC responses can "
            "exceed the permitted UDP message size (**1,232 bytes** [3, 11]), potentially causing **truncation**, "
            "**IP fragmentation**, **packet loss**, or **TCP fallback**.", 16)
    flow(s, MARGIN, 2.75, 2.15, 1.05, ["Larger PQC signatures and keys", "More transmitted bytes",
                                       "Potential latency, delivery and resource costs"], gap=0.42, size=14.5)
    textbox(s, MARGIN, 3.95, 7.3, 0.6, "A truncated answer must be asked for again over TCP: an extra round "
            "trip, more bytes and more work for every resolver.", 12.5, GREY)
    textbox(s, 8.25, 2.3, 4.5, 0.35, "Signature size vs one 1,232-byte UDP packet", 13, BLUE, bold=True)
    table(s, 8.25, 2.65, [2.2, 1.0, 1.3], [
        ["Signature type", "Size", "Fits?"],
        ["ECDSA P-256 (today)", "64 B", "yes"],
        ["RSA-2048 (today)", "256 B", "yes"],
        ["**SLH-DSA-128s (PQC, we test)**", "**7,856 B**", "**no, 6× too big**"],
    ], size=12, row_h=0.42)

    box(s, MARGIN, 4.6, WIDTH, 2.35, "", fill=PALE, line=BLUE, lw=1.5)
    textbox(s, MARGIN + 0.25, 4.72, WIDTH - 0.5, 0.35, "Core problem to investigate", 16, BLUE, bold=True)
    textbox(s, MARGIN + 0.25, 5.1, WIDTH - 0.5, 0.7, "**Can quantum-resistant DNSSEC authentication remain "
            "practical without excessive transmission and computational overhead?**", 17)
    textbox(s, MARGIN + 0.25, 5.85, WIDTH - 0.5, 0.35, "The trade-offs must be measured when comparing MTL and "
            "non-MTL workloads:", 14)
    for k, metric in enumerate(["Signing time", "Verification time", "Transmitted bytes", "Cache usage"]):
        box(s, MARGIN + 0.25 + k * 2.95, 6.3, 2.7, 0.45, [(f"**{metric}**", 14, BLUE)], fill=WHITE,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def background_size_slide(d: Deck):
    s = d.frame("Background", "Background (1/3): why size matters")
    textbox(s, MARGIN, 1.4, 5.9, 5.2, [
        "**DNSSEC** attaches a signature to each set of DNS records, so a resolver can check that an answer "
        "is genuine [1].",
        "Answers normally travel in **one small packet** (≤ 1,232 bytes) [3, 11].",
        "If an answer is bigger, it is **cut off**, and the resolver must ask again over a slower connection "
        "(TCP).",
        "**SLH-DSA** is the safest quantum-resistant option: it relies only on hash functions and has a tiny "
        "key. But its signature is **huge** [4].",
    ], 18, space=14, bullets=True)

    x0, label_w, scale = 6.95, 2.3, 3.3 / 7856
    textbox(s, x0, 1.45, 5.8, 0.35, "Signature size compared with one packet", 16, BLUE, bold=True)
    rows = [("ECDSA (today)", 64), ("RSA-2048", 256), ("SLH-DSA-128s (we test)", 7856)]
    bar_x, top = x0 + label_w, 2.2
    for k, (name, size) in enumerate(rows):
        y = top + k * 1.1
        fits = size <= 1232
        verdict = "fits" if fits else f"{size / 1232:.0f}× too big" if size > 2 * 1232 else "too big"
        textbox(s, x0, y - 0.05, label_w - 0.1, 0.3, f"**{name}**" if k == len(rows) - 1 else name, 14)
        textbox(s, x0, y + 0.27, label_w - 0.1, 0.3, f"**{size:,} B · {verdict}**", 12.5,
                GREEN if fits else "C62828")
        box(s, bar_x, y, max(size * scale, 0.05), 0.5, "", fill=GREEN if fits else "C62828", line=None)
    limit_x = bar_x + 1232 * scale
    arrow(s, limit_x, top - 0.2, limit_x, top + 3.05, BLACK, dash=True, head=False)
    textbox(s, limit_x - 1.0, top + 3.1, 2.0, 0.5, ["one packet", "1,232 B"], 12, align=PP_ALIGN.CENTER)


def background_mtl_slide(d: Deck):
    s = d.frame("Background", "Background (2/3): how MTL mode works")
    lw, x0, ly = 0.95, 0.7, 5.35
    leaf_cx = [x0 + i * 1.02 + lw / 2 for i in range(6)]
    for i, cx in enumerate(leaf_cx):
        box(s, cx - lw / 2, ly, lw, 0.5, f"RRset {i + 1}", fill=WHITE, line=GREY, size=11,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    node_y, rung_y, lad_y = 4.4, 3.4, 1.75
    nodes = [(leaf_cx[0] + leaf_cx[1]) / 2, (leaf_cx[2] + leaf_cx[3]) / 2]
    for k, cx in enumerate(nodes):
        box(s, cx - 0.17, node_y, 0.34, 0.34, "", fill=LIGHT, line=BLUE, shape=MSO_SHAPE.OVAL)
        for c in leaf_cx[2 * k: 2 * k + 2]:
            arrow(s, c, ly, cx, node_y + 0.34, BLUE, head=False)
    rungs = [((nodes[0] + nodes[1]) / 2, "rung 1"), ((leaf_cx[4] + leaf_cx[5]) / 2, "rung 2")]
    for cx, name in rungs:
        box(s, cx - 0.55, rung_y, 1.1, 0.5, name, fill="E2F2E3", line=GREEN, size=12,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    for cx in nodes:
        arrow(s, cx, node_y, rungs[0][0], rung_y + 0.5, BLUE, head=False)
    for c in leaf_cx[4:]:
        arrow(s, c, ly, rungs[1][0], rung_y + 0.5, BLUE, head=False)
    lad_x, lad_w = leaf_cx[0] - 0.2, leaf_cx[5] - leaf_cx[0] + 0.4
    box(s, lad_x, lad_y, lad_w, 0.85, [("**Ladder = {rung 1, rung 2}**", 14),
                                       ("signed **once** with SLH-DSA (≈ 8 KB)", 12)],
        fill="FFF3E0", line=AMBER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    for cx, _ in rungs:
        arrow(s, cx, rung_y, cx, lad_y + 0.85, AMBER, dash=True, head=False)

    rx, rw = 7.15, 5.6
    box(s, rx, 1.75, rw, 1.45, [("**Analogy:** instead of a notary signing every page of a long contract, the "
                                 "notary signs the **table of contents** once. Each page carries a short note "
                                 "proving where it sits in that table.", 15)],
        fill=PALE, line=None, align=PP_ALIGN.LEFT)
    table(s, rx, 3.45, [1.6, 4.0], [
        ["Term", "What it means"],
        ["**Ladder**", "the signed “table of contents” (≈ 8 KB), signed **once**"],
        ["**Short proof**", "what each record carries: ≈ 200 B instead of 7,856 B"],
        ["**Full signature**", "short proof + ladder; only needed if the resolver hasn't saved the ladder yet"],
        ["**Node hash**", "the quick step linking a record to the ladder; far cheaper than a signature"],
    ], size=13.5, row_h=0.62)
    textbox(s, MARGIN, 6.25, 6.4, 0.5, "MTL mode is defined in [5]; its use in DNSSEC in [6].", 12, GREY)


def lane(slide, x, title, messages, footer):
    """One message-sequence diagram: resolver on the left, authoritative server on the right."""
    left, right, top = x + 0.6, x + 3.4, 2.15
    textbox(slide, x, 1.45, 4.0, 0.35, title, 15, align=PP_ALIGN.CENTER, bold=True)
    box(slide, left - 0.55, 1.85, 1.1, 0.38, "Resolver", fill=PALE, size=11)
    box(slide, right - 0.6, 1.85, 1.2, 0.38, "Auth. server", fill=PALE, size=11)
    for lx in (left, right):
        arrow(slide, lx, top + 0.08, lx, 5.3, GREY, head=False, width=0.75)
    drop = 0.14
    for k, (text, to_server, color) in enumerate(messages):
        y = top + 0.45 + k * 0.5
        textbox(slide, left, y - 0.2, right - left, 0.22, text, 9.5, align=PP_ALIGN.CENTER)
        arrow(slide, left if to_server else right, y, right if to_server else left, y + drop, color, width=1.75)
    textbox(slide, x, 5.4, 4.0, 0.35, footer, 13.5, align=PP_ALIGN.CENTER, bold=True)


def background_lookup_slide(d: Deck):
    s = d.frame("Background", "Background (3/3): three ways a lookup can go")
    red = "C62828"
    lane(s, MARGIN - 0.1, "Plain SLH-DSA", [
        ("UDP query", True, BLUE), ("TC=1 (answer ≈ 8 KB does not fit)", False, red),
        ("TCP SYN", True, GREY), ("SYN-ACK", False, GREY), ("TCP query", True, BLUE),
        ("answer + 7,856 B RRSIG", False, red)], "3 round trips · ≈ 8 KB")
    lane(s, MARGIN + 4.05, "MTL, ladder cached (warm)", [
        ("UDP query", True, BLUE), ("answer + short proof (≈ 200 B)", False, GREEN)], "1 round trip · < 1,232 B")
    lane(s, MARGIN + 8.2, "MTL, no ladder (cold)", [
        ("UDP query", True, BLUE), ("short proof; ladder not known yet", False, AMBER),
        ("TCP SYN", True, GREY), ("SYN-ACK", False, GREY), ("TCP query + mtl-mode-full", True, BLUE),
        ("full signature (≈ 8.3 KB) → save ladder", False, AMBER)], "3 round trips once, then warm")
    box(s, MARGIN, 5.95, WIDTH, 0.85, [("**Key point:** MTL's benefit depends on **how big the zone is** and "
                                        "**how often the resolver can reuse a saved ladder**. Our test is built "
                                        "around these two factors.", 16)],
        fill="FFF3E0", line=None, align=PP_ALIGN.LEFT)


def literature_slide(d: Deck):
    s = d.frame("Literature & RFC Review", "Literature & RFC Review",
                "Only the sources that define what we build or measure.")
    rows = [
        ["Ref", "Source", "What it defines", "Used in Cipherops for"],
        ["[1]", "RFC 4033 / 4034 / 4035: DNSSEC", "RRSIG and DNSKEY formats; signing input (RFC 4034 §3.1.8.1)",
         "Byte-exact signing input per RRset (workload.py)"],
        ["[2]", "RFC 5155 NSEC3; RFC 9276", "Hashed denial of existence; 0 iterations, no salt",
         "NSEC3 chain in every zone; NXDOMAIN answers"],
        ["[3]", "RFC 6891 EDNS(0); RFC 7766", "UDP buffer size, TC bit, TCP fallback",
         "1,232 / 4,096-byte limits and TCP retries (dnswire.py)"],
        ["[4]", "NIST FIPS 205: SLH-DSA", "The PQC signature scheme (SLH-DSA-SHA2-128s)",
         "The algorithm under test, with and without MTL"],
        ["[5]", "draft-harvey-cfrg-mtl-mode-09", "MTL construction; §12 maximum sizes",
         "Analytical size and cache estimates (analytical.py)"],
        ["[6]", "draft-fregly-dnsop-slh-dsa-mtl-dnssec-05", "MTL in RRSIG: MTL-Type byte, mtl-mode-full option",
         "Encoding of MTL signatures in DNS answers"],
        ["[7]", "Fregly et al., CT-RSA 2023", "Design and size argument for MTL mode",
         "The hypothesis this benchmark tests"],
        ["[8]", "Verisign MTL reference implementation", "Supplied library, test suite and CLI tools",
         "MTL signing and verification; go/no-go gate"],
        ["[9]", "liboqs 0.16.0; OpenSSL 3.6.0", "Direct SLH-DSA; RSA, ECDSA, Ed25519",
         "Non-MTL and classical baselines"],
        ["[10]", "RFC 5702 / 6605 / 8080", "DNSSEC algorithms 8, 13 and 15", "Classical reference modes"],
        ["[11]", "DNS Flag Day 2020", "1,232-byte EDNS buffer recommendation", "Default UDP limit"],
    ]
    table(s, MARGIN, 1.62, [0.65, 3.55, 4.25, 3.7], rows, size=12.5, row_h=0.43)


def solution_slide(d: Deck):
    s = d.frame("Proposed Solution", "Proposed Solution")
    crop_top = 0.11  # the image's own title and subtitle duplicate the slide title
    top, bottom = 1.1, 6.98
    h = bottom - top
    w = h * (1672 / 941) / (1 - crop_top)
    pic = s.shapes.add_picture(str(ARCH), Inches((13.333 - w) / 2), Inches(top), Inches(w), Inches(h))
    pic.crop_top = crop_top


def optimization_slide(d: Deck):
    s = d.frame("Solution Optimization", "Optimization Proposed",
                "One SLH-DSA signature per zone instead of one per RRset. Large zone (10,005 RRsets), "
                "medians of 5 trials:")
    rows = [
        ["Metric", "SLH-DSA direct (before)", "SLH-DSA + MTL (after)", "Ed25519 (classical reference)"],
        ["Signature in each RRSIG", "7,856 B", "≤ 284 B condensed\n+ one 8,120 B signed ladder per zone", "64 B"],
        ["Sign the whole zone", "≈ 92 min (extrapolated)¹", "**0.58 s**", "0.37 s"],
        ["Verify one answer, ladder cached", "541 µs", "**7.2 µs**²", "117 µs"],
        ["Verify the first answer, nothing cached", "480 µs", "490 µs", "103 µs"],
        ["Signature storage for the zone", "78.6 MB", "**2.7 MB** (29× less)", "0.64 MB"],
        ["Answers that fit one 1,232-byte UDP packet", "0% (median 7,982 B)", "**87.4%** (median 411 B)",
         "91.5% (median 190 B)"],
        ["Bytes on the wire, 762 lookups³", "11.08 MB; 762 TCP retries", "**0.92 MB** (12× less); 246 TCP",
         "0.38 MB; 43 TCP"],
        ["Time for 762 lookups, 50 ms round trip³", "125.6 s", "**63.5 s** (2× faster)", "43.0 s"],
    ]
    table(s, MARGIN, 1.85, [3.75, 2.75, 3.35, 2.3], rows, size=13, row_h=0.43)
    textbox(s, MARGIN, 6.2, WIDTH, 0.8, [
        "¹ 1,000 of 10,005 RRsets signed per trial at 0.55 s each, scaled to the zone. Fully measured on the "
        "1,005-RRset zone: 558 s direct vs 0.57 s with MTL.   ² Includes one ladder verification spread over "
        "the zone.",
        "³ Measured live: real client and NSD server, EDNS 1,232 B, 2,000-lookup workload with resolver cache. "
        "MTL: 245 truncated answers plus 1 ladder fetch.",
    ], 11.5, GREY, space=2)


def uml_slide(d: Deck):
    s = d.frame("Architecture Design (UML Diagram)", "Solution Architecture & Design",
                "UML component view: one pinned container runs every stage and calls the unmodified "
                "reference libraries.")
    st = 8.5

    box(s, 0.55, 1.62, 2.65, 4.6, [("«device»  Windows host", 11)], fill=PALE, anchor=MSO_ANCHOR.TOP)
    box(s, 0.7, 2.1, 2.35, 0.9, [("«script»", st, GREY), ("**run_benchmark.ps1**", 11),
                                  ("-Config  -Cpuset  -Command", st), ("→ docker compose run", st, GREY)])
    box(s, 0.7, 3.35, 2.35, 1.0, [("«artifact»  volume mount", st, GREY), ("**Run results folder**", 11),
                                   ("config, environment, gate, tables, report, packet captures", st)])
    box(s, 0.7, 4.85, 2.35, 0.9, [("«component»", st, GREY), ("**Streamlit dashboard**", 11),
                                   ("charts, tables, packet viewer", st)])
    arrow(s, 1.875, 4.85, 1.875, 4.35)
    textbox(s, 1.95, 4.47, 1.1, 0.25, "reads", st, GREY)

    box(s, 3.5, 1.62, 9.3, 4.6, [("«execution environment»  Docker container  cipherops-bench:pinned", 11),
                                 ("Fedora 42 · OpenSSL 3.6.0 · liboqs 0.16.0 · Verisign MTL @7cef019 · "
                                  "1 CPU core · 4 GB · no network", st)], fill=PALE, anchor=MSO_ANCHOR.TOP)
    box(s, 3.7, 2.2, 8.9, 0.5, [("«component»  **runner.py**   main() → cmd_run(): stages 1–7", 11)])
    arrow(s, 3.05, 2.5, 3.7, 2.45)

    comps = [("preflight.py", "go / no-go"), ("workload.py", "zones, signing inputs"),
             ("harness.py", "run_trial()"), ("dnswire.py", "DNS answers, bytes"),
             ("cache.py", "analytical.py"), ("stats.py", "median, 95% CI"), ("report.py", "md, html, figures")]
    w, gap = (8.9 - 6 * 0.1) / 7, 0.1
    centers = []
    for i, (name, note) in enumerate(comps):
        x = 3.7 + i * (w + gap)
        centers.append(x + w / 2)
        box(s, x, 3.0, w, 0.85, [("«component»", st, GREY), (f"**{name}**", 10.5), (note, st)])
        arrow(s, x + w / 2, 2.7, x + w / 2, 3.0)

    hx = centers[2]
    box(s, hx - 0.8, 4.15, 1.6, 0.6, [("«executable»", st, GREY), ("**mtlbench (C)**", 10.5)])
    arrow(s, hx, 3.85, hx, 4.15)
    libs = [(3.7, "libmtlslib", "Verisign MTL"), (5.85, "liboqs", "SLH-DSA direct"),
            (8.0, "libcrypto", "OpenSSL: RSA, ECDSA, Ed25519")]
    for x, name, note in libs:
        box(s, x, 5.15, 2.0, 0.75, [("«library»", st, GREY), (f"**{name}**", 10.5), (note, st)], fill=LIGHT)
        arrow(s, hx, 4.75, x + 1.0, 5.15, dash=True)
    textbox(s, hx + 0.15, 4.8, 0.6, 0.25, "«use»", st, GREY)
    arrow(s, centers[0], 3.85, 4.4, 5.15, dash=True)
    textbox(s, 4.48, 4.2, 1.25, 0.45, ["«use»", "mtltest, CLI tools"], st, GREY)

    box(s, 10.25, 4.15, 2.35, 1.75, [("«component»  live validation", st, GREY),
                                      ("**live.py · crosscheck.py**", 10.5),
                                      ("real NSD server and client", st),
                                      ("LAN / WAN / lossy network profiles", st),
                                      ("tcpdump packet capture", st)], dash=True)
    arrow(s, 3.5, 3.85, 3.05, 3.85, dash=True)
    d.caption(s, "System Architecture UML Diagram")


def store(slide, x, y, w, h, ident, text, size=10):
    box(slide, x, y, w, h, [("", 1)], fill=PALE, line=None)
    for x1, y1, x2, y2 in ((x, y, x + w, y), (x, y + h, x + w, y + h), (x, y, x, y + h),
                           (x + 0.42, y, x + 0.42, y + h)):
        arrow(slide, x1, y1, x2, y2, color=BLUE, head=False)
    if ident:
        textbox(slide, x, y, 0.42, h, f"**{ident}**", size, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        textbox(slide, x + 0.5, y, w - 0.55, h, text.split("\n"), size, anchor=MSO_ANCHOR.MIDDLE)


def dfd_slide(d: Deck):
    s = d.frame("Data Flow Diagram", "Data Flow Diagram",
                "Level-1 DFD of one benchmark run: from the configuration to the report.")
    lab = 8.5
    box(s, 0.6, 1.62, 1.9, 0.6, [("**Operator**", 11), ("run_benchmark.ps1", lab)], fill=LIGHT)
    box(s, 4.4, 1.62, 4.3, 0.6, [("**Reference libraries**", 11), ("Verisign MTL · liboqs · OpenSSL", lab)],
        fill=LIGHT)
    box(s, 10.6, 1.62, 2.1, 0.6, [("**Reviewer / mentor**", 11)], fill=LIGHT)

    procs = [("1  Validate reference", "preflight.py"), ("2  Build zones", "workload.py"),
             ("3  Sign & verify", "harness.py → mtlbench"), ("4  DNS bytes & cache", "dnswire.py, cache.py"),
             ("5  Statistics & report", "stats.py, report.py")]
    cy, pw, ph = 3.3, 1.95, 1.05
    cx = [1.55 + i * 2.5 for i in range(5)]
    for (name, files), x in zip(procs, cx):
        box(s, x - pw / 2, cy - ph / 2, pw, ph, [(f"**{name}**", 11), (files, lab)], fill=WHITE,
            shape=MSO_SHAPE.OVAL)
    for i in range(4):
        arrow(s, cx[i] + pw / 2, cy, cx[i + 1] - pw / 2, cy)
    flow = ["GO", "", "", "tables"]
    for i, t in enumerate(flow):
        if t:
            textbox(s, cx[i] + pw / 2, cy - 0.3, 0.55, 0.25, t, lab, GREY, PP_ALIGN.CENTER)

    arrow(s, 1.55, 2.22, 1.55, cy - ph / 2)
    textbox(s, 1.65, 2.3, 1.2, 0.25, "start run", lab, GREY)
    arrow(s, 2.25, cy - 0.38, 4.4, 2.05, both=True)
    textbox(s, 2.6, 1.78, 1.6, 0.35, ["mtltest, tool round trip", "pass or fail"], lab, GREY, PP_ALIGN.CENTER)
    arrow(s, cx[2], 2.22, cx[2], cy - ph / 2, both=True)
    textbox(s, cx[2] + 0.1, 2.3, 2.3, 0.25, "signing inputs / signatures, timings", lab, GREY)
    arrow(s, cx[4], cy - ph / 2, cx[4], 2.22)
    textbox(s, cx[4] + 0.08, 2.3, 1.3, 0.45, "report, figures, dashboard", lab, GREY)

    sy, sh = 4.9, 0.62
    store(s, 0.45, sy, 2.2, sh, "D1", "Configuration\nseed, sizes, modes")
    store(s, 4.2, sy, 2.2, sh, "D2", "Zones\nzone files, signing inputs")
    store(s, 6.7, sy, 2.2, sh, "D3", "Raw measurements\ntimings, signatures")
    store(s, 9.2, sy, 2.2, sh, "D4", "Result tables\ntrials, bytes, cache")
    arrow(s, 2.65, sy + 0.1, cx[1] - 0.55, cy + 0.42)
    arrow(s, cx[1] + 0.35, cy + ph / 2 - 0.02, 4.9, sy)
    arrow(s, 5.9, sy, cx[2] - 0.35, cy + ph / 2 - 0.02)
    arrow(s, cx[2] + 0.35, cy + ph / 2 - 0.02, 7.4, sy)
    arrow(s, 8.4, sy, cx[3] - 0.35, cy + ph / 2 - 0.02)
    arrow(s, cx[3] + 0.35, cy + ph / 2 - 0.02, 9.9, sy)
    arrow(s, 10.9, sy, cx[4] - 0.35, cy + ph / 2 - 0.02)

    ly = 5.85
    box(s, 0.6, ly, 0.45, 0.25, "", fill=LIGHT)
    textbox(s, 1.12, ly, 1.6, 0.25, "External entity", 10)
    box(s, 2.75, ly, 0.45, 0.25, "", shape=MSO_SHAPE.OVAL)
    textbox(s, 3.27, ly, 1.0, 0.25, "Process", 10)
    store(s, 4.3, ly, 0.75, 0.25, "", "", 8)
    textbox(s, 5.12, ly, 1.1, 0.25, "Data store", 10)
    arrow(s, 6.25, ly + 0.125, 6.75, ly + 0.125)
    textbox(s, 6.85, ly, 1.0, 0.25, "Data flow", 10)
    d.caption(s, "Data Flow Diagram (DFD)")


def timeline_slide(d: Deck):
    s = d.frame("Deployment Timeline", "Deployment Timeline",
                "Milestones, and how the platform is deployed and reproduced.")
    steps = [
        ("Sep 30, 2026", "Scope and design", "PS-036 analysis, RFC and draft review, repository created", True),
        ("Oct 6, 2026", "Platform built", "Pinned image, go/no-go gate, C harness, zone builder, DNS and "
                                          "cache models, report, dashboard", True),
        ("Oct 7, 2026", "Quick demo run", "3 trials × 3 zones × 5 modes; gate GO; every correctness check "
                                          "passed", True),
        ("Oct 10–11, 2026", "Full PS-036 run", "5 trials; 50 / 500 / 5,000 names; gate GO; live validation "
                                               "on a real NSD server with packet capture", True),
        ("Next", "Submission", "Report; private GitHub repository transferred to aiori-hackathon; "
                               "mentor demo", False),
    ]
    w, gap, x0, ly = 2.27, 0.2, MARGIN, 1.9
    arrow(s, x0 + 0.2, ly, x0 + 5 * w + 4 * gap - 0.2, ly, color=BLUE, head=False, width=2.5)
    for i, (when, what, detail, done) in enumerate(steps):
        x = x0 + i * (w + gap)
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + w / 2 - 0.13), Inches(ly - 0.13), Inches(0.26),
                                 Inches(0.26))
        dot.shadow.inherit = False
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(BLUE if done else WHITE)
        dot.line.color.rgb = rgb(BLUE)
        dot.line.width = Pt(2)
        box(s, x, 2.2, w, 2.05, [(f"**{when}**", 13, BLUE), (f"**{what}**", 14), (detail, 11.5),
                                  (f"**{'Done' if done else 'Planned'}**", 11.5, GREEN if done else AMBER)],
            fill=PALE, anchor=MSO_ANCHOR.TOP)
    textbox(s, MARGIN, 4.45, WIDTH, 0.4, "Deployment methodology", 17, BLUE, bold=True)
    textbox(s, MARGIN, 4.9, WIDTH, 2.1, [
        "**One artifact.** Everything runs from one Docker image built from the repository; the host needs "
        "only Docker and git.",
        "**One command.** A single script builds the image and runs the benchmark (configuration and CPU cores "
        "selectable); a second runs the live validation on the finished run.",
        "**Self-describing runs.** Each run keeps its configuration, environment record, gate results, raw "
        "data, signatures, packet captures and report, so every number can be traced and re-run.",
        "**Tested.** pytest covers workloads, DNS wire format, cache and formulas on the host; the harness and "
        "library gates run inside the container.",
    ], 14, space=5, bullets=True)


REFS = [
    "R. Arends, R. Austein, M. Larson, D. Massey, S. Rose. “DNS Security Introduction and Requirements”; "
    "“Resource Records for the DNS Security Extensions”; “Protocol Modifications for the DNS Security "
    "Extensions”. In: ~~RFC 4033, 4034, 4035~~ (2005).",
    "B. Laurie, G. Sisson, R. Arends, D. Blacka. “DNS Security (DNSSEC) Hashed Authenticated Denial of "
    "Existence”. In: ~~RFC 5155~~ (2008). W. Hardaker, V. Dukhovni. “Guidance for NSEC3 Parameter Settings”. "
    "In: ~~RFC 9276~~ (2022).",
    "J. Damas, M. Graff, P. Vixie. “Extension Mechanisms for DNS (EDNS(0))”. In: ~~RFC 6891~~ (2013). "
    "J. Dickinson, S. Dickinson, R. Bellis, A. Mankin, D. Wessels. “DNS Transport over TCP – Implementation "
    "Requirements”. In: ~~RFC 7766~~ (2016).",
    "National Institute of Standards and Technology. “Stateless Hash-Based Digital Signature Standard”. "
    "In: ~~FIPS 205~~ (2024).",
    "J. Harvey et al. “Merkle Tree Ladder (MTL) Mode Signatures”. In: ~~Internet-Draft "
    "draft-harvey-cfrg-mtl-mode-09~~, IRTF CFRG (work in progress).",
    "A. Fregly et al. “Stateless Hash-Based Signatures in Merkle Tree Ladder Mode (SLH-DSA-MTL) for DNSSEC”. "
    "In: ~~Internet-Draft draft-fregly-dnsop-slh-dsa-mtl-dnssec-05~~, IETF DNSOP (work in progress).",
    "A. Fregly, J. Harvey, B. S. Kaliski Jr., S. Sheth. “Merkle Tree Ladder Mode: Reducing the Size Impact "
    "of NIST PQC Signature Algorithms in Practice”. In: ~~Topics in Cryptology – CT-RSA 2023~~, LNCS 13871, "
    "Springer (2023).",
    "Verisign. “MTL Mode reference implementation”. github.com/verisign/MTL, commit 7cef019.",
    "Open Quantum Safe. “liboqs”, version 0.16.0. github.com/open-quantum-safe/liboqs. OpenSSL Project. "
    "“OpenSSL”, version 3.6.0. openssl.org.",
    "J. Jansen. “Use of SHA-2 Algorithms with RSA in DNSKEY and RRSIG Resource Records for DNSSEC”. "
    "In: ~~RFC 5702~~ (2009). P. Hoffman, W. Wijngaards. “Elliptic Curve Digital Signature Algorithm (DSA) "
    "for DNSSEC”. In: ~~RFC 6605~~ (2012). O. Surý, R. Edmonds. “Edwards-Curve Digital Security Algorithm "
    "(EdDSA) for DNSSEC”. In: ~~RFC 8080~~ (2017).",
    "DNS Flag Day. “DNS Flag Day 2020”. dnsflagday.net/2020.",
]


def bibliography_slides(d: Deck):
    for part, (lo, hi) in (("I", (0, 6)), ("II", (6, len(REFS)))):
        s = d.frame("Bibliography", f"Bibliography {part}", root=False)
        tb = textbox(s, MARGIN, 1.25, WIDTH, 5.7, "", 15)
        tf = tb.text_frame
        for k in range(lo, hi):
            p = tf.paragraphs[0] if k == lo else tf.add_paragraph()
            if k > lo:
                p.space_before = Pt(9)
            rich(p, f"[{k + 1}]", 15, BLUE)
            rich(p, "\t" + REFS[k], 15)
            hanging(p, 0.6, None)


def thanks_slide(d: Deck):
    s = d.blank()
    textbox(s, 0, 3.55, 13.333, 0.6, "Thank You for Your Attention!", 23, align=PP_ALIGN.CENTER)


def main() -> None:
    d = Deck()
    title_slide(d)
    problem_slide(d)
    challenge_slide(d)
    background_size_slide(d)
    background_mtl_slide(d)
    background_lookup_slide(d)
    literature_slide(d)
    solution_slide(d)
    optimization_slide(d)
    uml_slide(d)
    dfd_slide(d)
    timeline_slide(d)
    bibliography_slides(d)
    thanks_slide(d)
    d.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
