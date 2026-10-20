"""Build docs/presentation/Cipherops-walkthrough-speaker-notes.pdf: each slide on top, its speaker script below.

Needs the slides PDF exported by PowerPoint next to the .pptx, plus python-pptx and PyMuPDF.

Usage: python scripts/make_walkthrough_notes_pdf.py
"""

from pathlib import Path

import pymupdf
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "docs" / "presentation" / "Cipherops-walkthrough"
FONT = Path("C:/Windows/Fonts/calibri.ttf")
BOLD = Path("C:/Windows/Fonts/calibrib.ttf")
BLUE = (0, 0, 0.9)


def main() -> None:
    slides = pymupdf.open(BASE.with_suffix(".pdf"))
    scripts = [s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
               for s in Presentation(BASE.with_suffix(".pptx")).slides]
    out = pymupdf.open()
    w, h = pymupdf.paper_size("a4")
    m = 48
    for i, text in enumerate(scripts):
        page = out.new_page(width=w, height=h)
        page.insert_font(fontname="cal", fontfile=str(FONT))
        page.insert_font(fontname="calb", fontfile=str(BOLD))
        page.insert_text((m, m), f"Slide {i + 1} / {len(scripts)}", fontname="calb", fontsize=13, color=BLUE)
        sw = w - 2 * m
        sh = sw * slides[i].rect.height / slides[i].rect.width
        frame = pymupdf.Rect(m, m + 12, m + sw, m + 12 + sh)
        page.show_pdf_page(frame, slides, i)
        page.draw_rect(frame, color=(0.75, 0.75, 0.75), width=0.5)
        page.insert_text((m, frame.y1 + 30), "Say:", fontname="calb", fontsize=12, color=BLUE)
        rest = page.insert_textbox(pymupdf.Rect(m, frame.y1 + 38, w - m, h - m), text, fontname="cal",
                                   fontsize=12, lineheight=1.35)
        if rest < 0:
            raise SystemExit(f"slide {i + 1}: script does not fit ({rest:.0f} pt short)")
    path = BASE.with_name(BASE.name + "-speaker-notes.pdf")
    out.save(path, garbage=3, deflate=True)
    print(path)


if __name__ == "__main__":
    main()
