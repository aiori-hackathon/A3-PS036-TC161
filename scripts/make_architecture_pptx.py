"""Build docs/architecture.pptx: editable architecture and run-flow slides.

Usage: pip install python-pptx==1.0.2 && python scripts/make_architecture_pptx.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "docs" / "architecture.pptx"

NAVY = "1F3A5F"
GREY_TEXT = "4A5568"
STYLES = {
    "container": ("F4F6FA", "8A99B3", NAVY),
    "orchestrator": ("1F3A5F", "1F3A5F", "FFFFFF"),
    "step": ("2F6FB0", "2F6FB0", "FFFFFF"),
    "lib": ("E8F1FB", "2F6FB0", NAVY),
    "live": ("FFF4E5", "E08A00", "7A4A00"),
    "input": ("EEF7EE", "3C8D40", "1E4D21"),
    "output": ("F3EEF9", "7B52AB", "3F2560"),
    "band": ("FFFFFF", "8A99B3", NAVY),
    "note": ("FFFFFF", "C9D2E0", GREY_TEXT),
}


def box(slide, x, y, w, h, title, lines=(), style="step", title_pt=13, body_pt=10,
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    fill, line, text = STYLES[style]
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill)
    s.line.color.rgb = RGBColor.from_string(line)
    s.line.width = Pt(1.25)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(0.06))
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = title
    r.font.bold = True
    r.font.size = Pt(title_pt)
    r.font.color.rgb = RGBColor.from_string(text)
    for line_text in lines:
        p = tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line_text
        r.font.size = Pt(body_pt)
        r.font.color.rgb = RGBColor.from_string(text)
    return s


def arrow(slide, x1, y1, x2, y2, color=NAVY, both=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    if both:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def text(slide, x, y, w, h, value, pt=12, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return tb


def architecture_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.3, 0.15, 12.7, 0.5, "Cipherops: high-level architecture", pt=26, bold=True)
    text(s, 0.3, 0.68, 12.7, 0.35,
         "PS-036 · PQC DNSSEC MTL benchmark using the Verisign MTL reference implementation · IETF 126 Hackathon",
         pt=12, color=GREY_TEXT)

    box(s, 0.3, 1.2, 1.95, 2.25, "INPUTS",
        ["configs/*.yaml", "", "seed · trials", "5 modes", "3 zone sizes", "query workload"],
        style="input", anchor=MSO_ANCHOR.TOP)
    box(s, 0.3, 3.6, 1.95, 2.35, "PS-036 RULES",
        ["No new crypto", "Pinned environment", "3 dataset sizes", "Repeated trials",
         "Measured ≠ analytical", "Defer if reference fails"],
        style="input", body_pt=9, anchor=MSO_ANCHOR.TOP)
    arrow(s, 2.25, 2.3, 2.6, 2.3)

    box(s, 2.6, 1.2, 8.0, 4.75, "Pinned Docker container",
        ["Fedora 42 · OpenSSL 3.6.0 · liboqs 0.16.0 · Verisign MTL @7cef019 · 1 CPU core · no network"],
        style="container", title_pt=12, body_pt=9, anchor=MSO_ANCHOR.TOP)
    box(s, 2.8, 1.85, 7.6, 0.45, "Orchestrator: runner.py   (python -m cipherops.runner run)",
        style="orchestrator", title_pt=12)

    steps = [("1 Go / No-Go", "preflight.py"), ("2 Zone builder", "workload.py"),
             ("3 Crypto harness", "mtlbench.c"), ("4 DNS answers & bytes", "dnswire.py"),
             ("5 Cache usage", "cache.py analytical.py"), ("6 Statistics", "stats.py"),
             ("7 Report", "report.py")]
    w, gap, x0, y = 0.98, 0.12, 2.8, 2.5
    for i, (title, file) in enumerate(steps):
        x = x0 + i * (w + gap)
        box(s, x, y, w, 0.95, title, [file], style="step", title_pt=10, body_pt=8)
        arrow(s, x + w / 2, 2.3, x + w / 2, y)
        if i:
            arrow(s, x - gap, y + 0.475, x, y + 0.475, color="8A99B3")

    harness_x = x0 + 2 * (w + gap) + w / 2
    box(s, 2.8, 3.75, 7.6, 1.0, "", style="note")
    text(s, 2.9, 3.77, 7.4, 0.3, "Crypto libraries (unmodified, called by the harness)", pt=10, bold=True)
    libs = [("Verisign MTL library", "MTL mode: leaves, ladder, condensed sigs"),
            ("liboqs", "SLH-DSA-SHA2-128s (direct)"),
            ("OpenSSL", "RSA-2048 · ECDSA P-256 · Ed25519")]
    for i, (title, line) in enumerate(libs):
        box(s, 2.95 + i * 2.5, 4.07, 2.3, 0.6, title, [line], style="lib", title_pt=10, body_pt=8)
    arrow(s, harness_x, y + 0.95, harness_x, 3.75)

    box(s, 2.8, 4.9, 7.6, 0.95, "", style="note")
    text(s, 2.9, 4.92, 7.4, 0.3, "Real-traffic validation (live mode / cross-check)", pt=10, bold=True)
    box(s, 2.95, 5.2, 2.3, 0.55, "Live client", ["live.py: UDP → TC → TCP"], style="live", title_pt=10, body_pt=8)
    box(s, 5.45, 5.2, 2.3, 0.55, "NSD server", ["serves the signed zone"], style="live", title_pt=10, body_pt=8)
    box(s, 7.95, 5.2, 2.3, 0.55, "tcpdump", ["packets → .pcap"], style="live", title_pt=10, body_pt=8)
    arrow(s, 5.25, 5.475, 5.45, 5.475, color="E08A00", both=True)
    arrow(s, 7.75, 5.475, 7.95, 5.475, color="E08A00")

    arrow(s, 10.6, 2.3, 10.95, 2.3)
    box(s, 10.95, 1.2, 2.1, 2.6, "OUTPUTS",
        ["results/runs/<run_id>/", "", "trials.csv", "responses · workload", "cache · sizes", "summary.json",
         "report md/html/pdf", "live.csv · pcap"],
        style="output", body_pt=9, anchor=MSO_ANCHOR.TOP)
    box(s, 10.95, 4.0, 2.1, 0.85, "Dashboard", ["Streamlit (dashboard/app.py)"], style="output", title_pt=12, body_pt=9)
    box(s, 10.95, 5.0, 2.1, 0.85, "Wireshark", ["opens the .pcap files"], style="output", title_pt=12, body_pt=9)
    arrow(s, 12.0, 3.8, 12.0, 4.0, color="7B52AB")

    box(s, 0.3, 6.15, 12.75, 0.5, "Compared on identical zones:   RSA-2048  ·  ECDSA P-256  ·  Ed25519  ·  "
        "SLH-DSA-SHA2-128s  ·  SLH-DSA-SHA2-128s + MTL", style="band", title_pt=12)
    box(s, 0.3, 6.75, 12.75, 0.5, "Measured:   signing time  ·  verification time  ·  transmitted bytes "
        "(UDP 1232 / TCP fallback)  ·  cache usage", style="band", title_pt=12)


def flow_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.3, 0.15, 12.7, 0.5, "Cipherops: end-to-end run flow", pt=26, bold=True)
    box(s, 0.35, 0.85, 12.6, 0.5,
        ".\\scripts\\run_benchmark.ps1   →   docker compose run bench   →   python -m cipherops.runner run --config configs/quick.yaml",
        style="orchestrator", title_pt=12)

    steps = [
        ("1  Configure", "runner.py · envinfo.py", ["Load YAML", "Record CPU, versions, git SHA"],
         ["config.yaml", "environment.json"]),
        ("2  Go / No-Go", "preflight.py", ["Verisign's 89 tests", "Tools round trip", "Harness smoke test"],
         ["preflight.json", "GO → continue", "NO-GO → stop"]),
        ("3  Build zones", "workload.py", ["bench.example in 3 sizes", "RRsets = 2·names + 5",
                                          "RFC 4034 signing inputs"],
         ["zone.txt", "messages.bin", "manifest.csv"]),
        ("4  Sign & verify", "harness.py → mtlbench.c", ["Warm-up, then N trials", "Shuffled mode order",
                                                         "Tamper checks"],
         ["trials.csv", "ops.parquet", "artifacts/ (signatures)"]),
        ("5  Bytes & cache", "dnswire.py · cache.py", ["Real DNS answers", "UDP 1232 vs TC → TCP",
                                                       "Resolver cache memory"],
         ["responses.csv", "workload.csv", "cache.csv · sizes.csv"]),
        ("6  Statistics", "stats.py", ["Median per trial set", "Bootstrap 95% CI"], ["summary.json"]),
        ("7  Report", "report.py", ["Figures", "Data-driven text"], ["report.md / .html", "figures/", "→ dashboard"]),
    ]
    w, gap, x0 = 1.7, 0.12, 0.35
    for i, (title, file, does, makes) in enumerate(steps):
        x = x0 + i * (w + gap)
        box(s, x, 1.65, w, 0.85, title, [file], style="step", title_pt=13, body_pt=9)
        box(s, x, 2.65, w, 1.6, "What it does", ["• " + d for d in does], style="note", title_pt=11,
            body_pt=11, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT)
        box(s, x, 4.4, w, 1.35, "Produces", makes, style="output", title_pt=11, body_pt=11,
            anchor=MSO_ANCHOR.TOP)
        if i:
            arrow(s, x - gap, 2.075, x, 2.075)

    box(s, 0.35, 6.0, 12.6, 0.95, "Optional, after a run",
        ["runner live  → real NSD server + real client + tcpdump pcap        "
         "runner crosscheck  → model bytes vs NSD        streamlit run dashboard/app.py        "
         "scripts/make_pdf.ps1  → report.pdf"],
        style="live", title_pt=11, body_pt=10)


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    architecture_slide(prs)
    flow_slide(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
